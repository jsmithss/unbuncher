#!/usr/bin/env python3
"""
Excel Cell Unmerger - Standalone Executable Version
Unmerges cells in Excel files and fills all cells with the merged content
"""

import sys
import os
import time
from pathlib import Path

# For PyInstaller compatibility
if getattr(sys, 'frozen', False):
    # Running as compiled executable
    application_path = os.path.dirname(sys.executable)
else:
    # Running as script
    application_path = os.path.dirname(os.path.abspath(__file__))

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print("=" * 60)
    print("ERROR: Required Excel library not found!")
    print("This should not happen in the standalone version.")
    print(f"Technical details: {e}")
    print("=" * 60)
    input("Press Enter to exit...")
    sys.exit(1)

def unmerge_excel(input_file):
    """Unmerge cells in Excel file and fill with content"""
    
    # Check if file exists
    if not os.path.exists(input_file):
        print(f"Error: File not found: {input_file}")
        return False
    
    try:
        # Load workbook
        print(f"Loading workbook: {input_file}")
        workbook = openpyxl.load_workbook(input_file)
        
        # Process each worksheet
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            print(f"Processing sheet: {sheet_name}")
            
            # Get all merged cell ranges (make a copy since we'll modify the original)
            merged_ranges = list(worksheet.merged_cells.ranges)
            
            # Unmerge cells and fill with content
            for merged_range in merged_ranges:
                # Get the value from the top-left cell of the merged range
                top_left_cell = worksheet.cell(merged_range.min_row, merged_range.min_col)
                merged_value = top_left_cell.value
                
                # Unmerge the cells
                worksheet.unmerge_cells(str(merged_range))
                
                # Fill all cells in the range with the merged value
                for row in range(merged_range.min_row, merged_range.max_row + 1):
                    for col in range(merged_range.min_col, merged_range.max_col + 1):
                        worksheet.cell(row, col).value = merged_value
        
        # Create output filename
        input_path = Path(input_file)
        output_file = input_path.parent / f"{input_path.stem}_nomerge{input_path.suffix}"
        
        # Save the workbook
        workbook.save(output_file)
        print(f"Saved unmerged data to: {output_file}")
        return True
        
    except Exception as e:
        print(f"Error processing file: {e}")
        return False

def main():
    print("=" * 60)
    print("Excel Cell Unmerger v1.0")
    print("=" * 60)
    
    if len(sys.argv) != 2:
        print("ERROR: No file provided!")
        print("\nUsage:")
        print("  • Drag and drop an Excel file onto this executable")
        print("  • Or run from command line: unmerge_excel.exe <filename.xlsx>")
        print("\nSupported formats: .xlsx, .xlsm")
        input("\nPress Enter to exit...")
        sys.exit(1)
    
    input_file = sys.argv[1]
    
    # Validate file extension
    valid_extensions = ['.xlsx', '.xlsm']
    file_ext = Path(input_file).suffix.lower()
    
    if file_ext not in valid_extensions:
        print(f"ERROR: Unsupported file format '{file_ext}'")
        print(f"Supported formats: {', '.join(valid_extensions)}")
        input("\nPress Enter to exit...")
        sys.exit(1)
    
    print(f"Processing: {os.path.basename(input_file)}")
    print("-" * 40)
    
    start_time = time.time()
    
    if unmerge_excel(input_file):
        elapsed_time = time.time() - start_time
        print("-" * 40)
        print(f"✓ Processing completed successfully in {elapsed_time:.1f} seconds!")
        print("\nYou can now open the '_nomerge' file in Excel.")
    else:
        print("-" * 40)
        print("✗ Processing failed!")
        input("\nPress Enter to exit...")
        sys.exit(1)
    
    print("\nPress Enter to exit...")
    input()

if __name__ == "__main__":
    main()
